"""
국가별 Word 양식(.docx) 템플릿 엔진.

양식 파일 안에 {{모델명}} 같은 자리표시자를 쓰면, 문서 생성 시 입력값으로 치환한다.
사이트의 "양식 관리" 화면에서 업로드한 파일은 templates/custom/<국가>.docx 로 저장되고,
파일이 있으면 그 양식으로, 없으면 기존 코드 빌더(doc_*.py)로 문서를 만든다.
"""
import re
from io import BytesIO
from pathlib import Path

from docx import Document

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates" / "custom"

COUNTRIES = ["FDA", "HC", "EU", "KR"]

# 모든 국가에서 쓸 수 있는 자리표시자 (키 -> 설명)
COMMON_PLACEHOLDERS = {
    "모델명": "제품 모델명",
    "의료기기등급": "의료기기 등급 (Class I ~ III)",
    "제조사": "제조사(Legal Manufacturer)",
    "제조사주소": "제조사 주소",
    "FDA_510k번호": "FDA 510(k) 번호",
    "FDA_제품코드": "FDA Product Code",
    "FDA_규정번호": "FDA Regulation Number",
    "HC_라이선스번호": "Health Canada Licence No.",
    "EU_인증서번호": "EU NB 인증서 번호",
    "EU_인증기관": "EU Notified Body",
    "EU_BasicUDIDI": "EU Basic UDI-DI",
    "국내허가번호": "국내 허가·인증·신고번호",
    "구성품명": "변경 대상 구성품",
    "변경요약": "변경 사항 한줄 요약",
    "변경전": "변경 전(Before)",
    "변경후": "변경 후(After)",
    "변경사유": "변경 사유",
    "변경상세": "변경 상세 설명 (없으면 한줄 요약)",
    "평가일": "평가일/발효일",
    "문서번호": "문서 번호",
    "개정번호": "Revision No.",
    "작성자": "작성자(Prepared by)",
    "검토자": "검토자(Reviewed by)",
    "승인자": "승인자(Approved by)",
    "작성일": "작성일",
    "검토일": "검토일",
    "승인일": "승인일",
    "설계사양서번호": "Design Specifications 문서번호",
    "위험평가서번호": "Risk-Based Assessment 문서번호",
    "VV요약서번호": "Verification and Validation Summary 문서번호",
    "가이던스명": "적용한 가이던스 이름",
    "판정결과": "평가 결과 (조치 사항)",
    "판정요약": "평가 결과 요약 문장",
    "판단흐름": "질문과 답변 목록 (여러 줄로 자동 확장)",
}

# FDA 양식의 체크박스용
FDA_PLACEHOLDERS = {
    "체크_510k제출": "Submit 510(k) 체크박스 ([X] 또는 [ ])",
    "체크_LetterToFile": "Letter to file 체크박스 ([X] 또는 [ ])",
}

_TOKEN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def placeholders_for(country: str) -> dict:
    result = dict(COMMON_PLACEHOLDERS)
    if country == "FDA":
        result.update(FDA_PLACEHOLDERS)
    return result


def template_path(country: str) -> Path:
    return TEMPLATE_DIR / f"{country}.docx"


def has_custom_template(country: str) -> bool:
    return template_path(country).exists()


def save_custom_template(country: str, data: bytes) -> None:
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    template_path(country).write_bytes(data)


def delete_custom_template(country: str) -> None:
    p = template_path(country)
    if p.exists():
        p.unlink()


def _answer_text(country: str, answer) -> str:
    if country == "KR":
        return "예" if answer else "아니오"
    return "Yes" if answer else "No"


def _flow_lines(country: str, assessment: dict) -> str:
    lines = []
    for q in assessment.get("path", []):
        if q.get("answer") is None:
            continue
        text = q.get("text_en") if country in ("FDA", "HC") and q.get("text_en") else q.get("text", "")
        lines.append(f"{q.get('id', '')}. {text} → {_answer_text(country, q['answer'])}")
    return "\n".join(lines)


def build_context(country: str, product_info: dict, change_info: dict,
                  assessment: dict, metadata: dict) -> dict:
    if assessment.get("manualReviewRequired"):
        submit_selected = False
    else:
        submit_selected = bool(assessment.get("isSignificant"))

    ctx = {
        "모델명": product_info.get("modelName", ""),
        "의료기기등급": product_info.get("deviceClass", ""),
        "제조사": product_info.get("manufacturer", ""),
        "제조사주소": product_info.get("manufacturerAddress", ""),
        "FDA_510k번호": product_info.get("fdaK510", ""),
        "FDA_제품코드": product_info.get("fdaProductCode", ""),
        "FDA_규정번호": product_info.get("fdaRegulationNumber", ""),
        "HC_라이선스번호": product_info.get("hcLicenceNo", ""),
        "EU_인증서번호": product_info.get("euCertNumber", ""),
        "EU_인증기관": product_info.get("euNotifiedBody", ""),
        "EU_BasicUDIDI": product_info.get("euBasicUDIDI", ""),
        "국내허가번호": product_info.get("krPermitNo", ""),
        "구성품명": change_info.get("componentName", ""),
        "변경요약": change_info.get("changeTitle", ""),
        "변경전": change_info.get("beforeValue", ""),
        "변경후": change_info.get("afterValue", ""),
        "변경사유": change_info.get("reason", ""),
        "변경상세": change_info.get("description") or change_info.get("changeTitle", ""),
        "평가일": metadata.get("effectiveDate", ""),
        "문서번호": metadata.get("docNumber", ""),
        "개정번호": metadata.get("revisionNo", ""),
        "작성자": metadata.get("preparedBy", ""),
        "검토자": metadata.get("reviewedBy", ""),
        "승인자": metadata.get("approvedBy", ""),
        "작성일": metadata.get("preparedDate", ""),
        "검토일": metadata.get("reviewedDate", ""),
        "승인일": metadata.get("approvedDate", ""),
        "설계사양서번호": change_info.get("designSpecRef", ""),
        "위험평가서번호": change_info.get("riskAssessmentRef", ""),
        "VV요약서번호": change_info.get("vvSummaryRef", ""),
        "가이던스명": assessment.get("guidance", ""),
        "판정결과": assessment.get("requiredAction", ""),
        "판정요약": assessment.get("summary", ""),
        "판단흐름": _flow_lines(country, assessment),
        "체크_510k제출": "[X]" if submit_selected else "[ ]",
        "체크_LetterToFile": "[ ]" if submit_selected else "[X]",
    }
    return {k: ("" if v is None else str(v)) for k, v in ctx.items()}


def _replace_in_text(text: str, context: dict) -> str:
    return _TOKEN.sub(lambda m: context.get(m.group(1), m.group(0)), text)


def _process_paragraph(paragraph, context: dict) -> None:
    if "{{" not in paragraph.text:
        return
    # 1) 자리표시자가 한 run 안에 온전히 있으면 run 단위로 치환 (서식 유지)
    for run in paragraph.runs:
        if "{{" in run.text:
            run.text = _replace_in_text(run.text, context)
    # 2) Word가 자리표시자를 여러 run으로 쪼갠 경우: 문단 전체를 합쳐 치환
    if "{{" in paragraph.text and paragraph.runs:
        merged = _replace_in_text(paragraph.text, context)
        paragraph.runs[0].text = merged
        for run in paragraph.runs[1:]:
            run.text = ""


def _process_container(container, context: dict) -> None:
    for p in container.paragraphs:
        _process_paragraph(p, context)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                _process_container(cell, context)


def render_docx(template_bytes: bytes, context: dict):
    doc = Document(BytesIO(template_bytes))
    _process_container(doc, context)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            _process_container(part, context)
    return doc


def sample_context(country: str) -> dict:
    """미리보기용 샘플 데이터로 채운 치환값"""
    product_info = {
        "modelName": "RIOScan", "deviceClass": "Class IIb", "manufacturer": "Ray Co., Ltd.",
        "manufacturerAddress": "Yongin-si, Gyeonggi-do, Republic of Korea", "fdaK510": "K123456",
        "fdaProductCode": "MQB", "fdaRegulationNumber": "892.1650", "hcLicenceNo": "12345",
        "euCertNumber": "NB-0123", "euNotifiedBody": "Example NB", "euBasicUDIDI": "BUDI-0001",
        "krPermitNo": "제허 00-000호",
    }
    change_info = {
        "componentName": "AC/DC power adapter", "changeTitle": "Power adapter model change",
        "beforeValue": "Model A", "afterValue": "Model B", "reason": "Original supplier discontinued Model A.",
        "description": "The AC/DC power adapter was changed from Model A to Model B.",
        "designSpecRef": "DS-0001", "riskAssessmentRef": "RA-0001", "vvSummaryRef": "VV-0001",
    }
    assessment = {
        "guidance": "Sample guidance title", "isSignificant": False, "requiredAction": "Not significant (sample)",
        "summary": "Sample conclusion: the change is not significant.",
        "path": [
            {"id": "Q1", "text": "샘플 질문 1", "text_en": "Sample question 1", "answer": False},
            {"id": "Q2", "text": "샘플 질문 2", "text_en": "Sample question 2", "answer": False},
        ],
    }
    metadata = {
        "effectiveDate": "2026-01-01", "docNumber": "DOC-2026-SAMPLE", "revisionNo": "00",
        "preparedBy": "홍길동", "reviewedBy": "김검토", "approvedBy": "이승인",
        "preparedDate": "2026-01-01", "reviewedDate": "2026-01-02", "approvedDate": "2026-01-03",
    }
    return build_context(country, product_info, change_info, assessment, metadata)


def _esc(text: str) -> str:
    import html
    s = html.escape(text).replace("\n", "<br>")
    # 치환되지 않고 남은 자리표시자는 노란색으로 표시
    return re.sub(r"\{\{(.+?)\}\}", r'<mark style="background:#ffe08a;color:#000">{{\1}}</mark>', s)


def _run_html(run) -> str:
    text = run.text
    if not text:
        return ""
    out = _esc(text)
    style = []
    try:
        if run.font.color is not None and run.font.color.rgb is not None:
            style.append(f"color:#{run.font.color.rgb}")
    except Exception:
        pass
    try:
        if run.font.size:
            style.append(f"font-size:{run.font.size.pt}pt")
    except Exception:
        pass
    if style:
        out = f'<span style="{";".join(style)}">{out}</span>'
    if run.bold:
        out = f"<b>{out}</b>"
    if run.italic:
        out = f"<i>{out}</i>"
    return out


def _para_html(paragraph) -> str:
    if paragraph._p.xpath('.//w:br[@w:type="page"]'):
        return '<div style="border-top:2px dashed #999;margin:18px 0;color:#999;font-size:11px">— 페이지 나눔 —</div>'
    inner = "".join(_run_html(r) for r in paragraph.runs)
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    a = paragraph.alignment
    align = {1: "center", 2: "right"}.get(int(a) if a is not None else 0, "left")
    if not inner.strip():
        return '<div style="height:10px"></div>'
    if style_name.startswith("Heading 1"):
        return f'<h2 style="margin:14px 0 6px;text-align:{align}">{inner}</h2>'
    if style_name.startswith("Heading"):
        return f'<h3 style="margin:10px 0 4px;text-align:{align}">{inner}</h3>'
    prefix = "• " if "List" in style_name else ""
    return f'<p style="margin:0 0 8px;text-align:{align}">{prefix}{inner}</p>'


def _cell_fill(cell) -> str:
    try:
        shd = cell._tc.xpath("./w:tcPr/w:shd")
        if shd:
            fill = shd[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill")
            if fill and fill.lower() not in ("auto", "ffffff"):
                return f"background:#{fill};"
    except Exception:
        pass
    return ""


def _table_html(table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        seen = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            body = _container_html(cell)
            cells.append(f'<td style="border:1px solid #888;padding:6px 8px;vertical-align:top;{_cell_fill(cell)}">{body}</td>')
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return '<table style="border-collapse:collapse;width:100%;margin:6px 0 12px">' + "".join(rows) + "</table>"


def _container_html(container) -> str:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    parts = []
    for child in container._element.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(_para_html(Paragraph(child, container)))
        elif tag == "tbl":
            parts.append(_table_html(Table(child, container)))
    return "".join(parts)


def docx_to_preview_html(doc) -> str:
    """python-docx Document를 종이 문서 모양의 HTML로 변환 (미리보기용, Word와 100% 동일하지는 않음)"""
    body = _container_html(doc._body)
    footer_texts = []
    for section in doc.sections:
        t = " ".join(p.text for p in section.footer.paragraphs if p.text.strip())
        if t:
            footer_texts.append(t)
    footer = ""
    if footer_texts:
        footer = ('<div style="border-top:1px solid #ccc;margin-top:18px;padding-top:6px;color:#666;font-size:11px">'
                  "바닥글: " + _esc(footer_texts[0]) + "</div>")
    return ('<div style="background:#fff;color:#000;border:1px solid #bbb;box-shadow:0 2px 8px rgba(0,0,0,.15);'
            'padding:32px 40px;font-family:Arial,sans-serif;font-size:11pt;line-height:1.45;overflow-x:auto">'
            + body + footer + "</div>")


def find_placeholders(template_bytes: bytes) -> set:
    """양식 안에서 발견된 자리표시자 키 집합"""
    found = set()

    def scan(container):
        for p in container.paragraphs:
            found.update(m.group(1) for m in _TOKEN.finditer(p.text))
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    scan(cell)

    doc = Document(BytesIO(template_bytes))
    scan(doc)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            scan(part)
    return found
