"""
국가별 "기본 양식 파일" 생성기.

현재 코드 빌더(doc_*.py)가 만드는 문서와 같은 구성으로, 입력값 자리에 {{자리표시자}}를 넣은
Word 파일을 만든다. 사이트의 "양식 관리"에서 내려받아 Word로 편집한 뒤 다시 올리면 된다.
"""
from io import BytesIO

from docx.enum.text import WD_ALIGN_PARAGRAPH

from doc_common import (
    init_document, add_paragraph, add_spacer, add_page_break, add_labeled_paragraph,
    add_info_table, add_before_after_table, add_signature_table, add_generic_table,
    add_cover_page, add_revision_history, add_toc, add_footer, add_h1, add_bullet,
)


def _fda():
    doc = init_document()
    add_paragraph(doc, "Regulatory Change Assessment", bold=True, size=16,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, after_pt=16)
    add_labeled_paragraph(doc, "Product Name", "{{모델명}}")
    add_labeled_paragraph(doc, "Date of Assessment", "{{평가일}}")
    add_labeled_paragraph(doc, "Device Description", "{{모델명}} ({{의료기기등급}}, {{제조사}})")
    add_labeled_paragraph(doc, "Description of Change(s)", "{{변경상세}}")
    add_labeled_paragraph(doc, "Reason for Change(s)", "{{변경사유}}")
    add_labeled_paragraph(
        doc,
        "Applicable Regulatory History (including 510(k) #s and comparison of modified device to most recently cleared version)",
        "Most recently cleared 510(k): {{FDA_510k번호}}.", after_pt=4,
    )
    add_before_after_table(doc, [("{{구성품명}}", "{{변경전}}", "{{변경후}}")])
    add_spacer(doc)
    add_paragraph(doc, "Completed Checklist Attached:", bold=True, after_pt=2)
    add_paragraph(doc, "[X] Yes", after_pt=2)
    add_paragraph(doc, "[ ] No (include rationale if selected)", after_pt=14)
    add_paragraph(doc, "Recommended Regulatory Action:", bold=True, after_pt=2)
    add_paragraph(doc, "{{체크_510k제출}} Submit 510(k)", after_pt=2)
    add_paragraph(doc, "{{체크_LetterToFile}} Letter to file", after_pt=14)
    add_paragraph(doc, "Supporting Documents:", bold=True, after_pt=2)
    add_paragraph(doc, "Design Specifications: {{설계사양서번호}}", after_pt=2)
    add_paragraph(doc, "Risk-Based Assessment: {{위험평가서번호}}", after_pt=2)
    add_paragraph(doc, "Verification and Validation Summary: {{VV요약서번호}}", after_pt=14)
    add_paragraph(doc, "Signatures", bold=True)
    add_signature_table(doc, [
        ("Prepared by", "{{작성자}}", "{{작성일}}"),
        ("Reviewed by", "{{검토자}}", "{{검토일}}"),
        ("Approved by", "{{승인자}}", "{{승인일}}"),
    ])
    add_page_break(doc)
    add_paragraph(doc, "Attachment: Completed Checklist", bold=True, size=14)
    add_paragraph(doc, "Assessed in accordance with FDA's guidance \"Deciding When to Submit a 510(k) for a Change to an Existing Device\".", italic=True)
    add_spacer(doc)
    add_paragraph(doc, "{{판단흐름}}")
    return doc


def _hc():
    doc = init_document()
    add_paragraph(doc, "{{평가일}}")
    add_paragraph(doc, "Health Canada", after_pt=0)
    add_paragraph(doc, "Medical Devices Bureau", after_pt=0)
    add_paragraph(doc, "Device Licensing Services Division", after_pt=0)
    add_paragraph(doc, "[Address]")
    add_spacer(doc)
    add_paragraph(doc, "Re: Notification of Non-Significant Device Change — {{모델명}} (Medical Device Licence No. {{HC_라이선스번호}})", bold=True)
    add_paragraph(doc, "To Whom It May Concern:")
    add_paragraph(doc, "Pursuant to Section 43(1)(b) of the Medical Devices Regulations (SOR/98-282), {{제조사}} hereby notifies Health Canada of a change made to the above-referenced licensed medical device, {{모델명}} (a {{의료기기등급}} medical device), as part of our next annual licence renewal filing.")
    add_paragraph(doc, "Description of Change: {{변경상세}} {{변경사유}}")
    add_paragraph(doc, "Basis for Determination: {{모델명}} is licensed as a {{의료기기등급}} medical device. This change was assessed in accordance with Health Canada's guidance document, {{가이던스명}}. Section 34 of the Regulations sets out the categories of changes, under paragraphs (a) through (f), that require submission of a licence amendment application; this change does not fall within any of these categories. Accordingly, a licence amendment is not required for this change.")
    add_paragraph(doc, "This change has been evaluated and documented within {{제조사}}'s quality management system (ISO 13485) under our design change control procedures, including verification and/or validation activities confirming that the safety and performance of the device are unaffected by this change.")
    add_paragraph(doc, "In accordance with Health Canada's guidance recommending that non-significant changes be itemized in a table with a brief rationale at the time of annual licence renewal, the details of this change are provided below:")
    add_spacer(doc)
    add_generic_table(
        doc,
        ["Date of Change", "Device / Change", "Rationale (Why Not Significant)", "Supporting Documentation"],
        [["{{평가일}}", "{{모델명}} — {{구성품명}}, {{변경전}} → {{변경후}}",
          "{{변경사유}}",
          "{{설계사양서번호}}; {{위험평가서번호}}; {{VV요약서번호}}"]],
        col_widths_cm=[2.7, 4.3, 6.5, 2.5],
    )
    add_spacer(doc)
    add_paragraph(doc, "Should you require any additional information regarding this change, please do not hesitate to contact the undersigned.")
    add_paragraph(doc, "Sincerely,")
    add_paragraph(doc, "{{작성자}}", after_pt=0)
    add_paragraph(doc, "[Title]", after_pt=0)
    add_paragraph(doc, "Regulatory Affairs, {{제조사}}", after_pt=0)
    add_paragraph(doc, "[Email] / [Phone]")
    return doc


def _eu():
    doc = init_document()
    add_footer(doc, "{{문서번호}}")
    add_cover_page(
        doc,
        doc_title="NON-SIGNIFICANT CHANGE ASSESSMENT",
        doc_subtitle="EU MDR (Regulation (EU) 2017/745) – per MDCG 2020-3\n{{변경요약}}",
        doc_number="{{문서번호}}",
        rev_no="{{개정번호}}",
        effective_date="{{평가일}}",
        prepared_by="{{작성자}}",
        reviewed_by="{{검토자}}",
        approved_by="{{승인자}}",
    )
    add_revision_history(doc, {"revisionNo": "{{개정번호}}", "effectiveDate": "{{평가일}}", "preparedBy": "{{작성자}}"})
    add_toc(doc)
    add_h1(doc, "1. Purpose & Scope")
    add_paragraph(doc, 'This document evaluates a proposed change to the device under the requirements of Regulation (EU) 2017/745 (MDR) and the guidance MDCG 2020-3 Rev.1 "Guidance on significant changes regarding the transitional provision under Article 120 of the MDR." The assessment determines whether the change qualifies as a significant change in design or intended purpose, which would require prior notification to and approval by the Notified Body.')
    add_h1(doc, "2. Device & Certificate Information")
    add_info_table(doc, [
        ("Model Name", "{{모델명}}"),
        ("Risk Class (MDR Annex VIII)", "{{의료기기등급}}"),
        ("Basic UDI-DI", "{{EU_BasicUDIDI}}"),
        ("Legal Manufacturer", "{{제조사}}"),
        ("Manufacturer Address", "{{제조사주소}}"),
        ("Notified Body", "{{EU_인증기관}}"),
        ("Certificate Number", "{{EU_인증서번호}}"),
    ])
    add_h1(doc, "3. Description of the Change")
    add_before_after_table(doc, [["{{구성품명}} – Item", "{{변경전}}", "{{변경후}}"]])
    add_spacer(doc)
    add_paragraph(doc, "Reason: {{변경사유}}")
    add_paragraph(doc, "Description: {{변경상세}}")
    add_h1(doc, "4. Assessment per MDCG 2020-3 Decision Charts")
    add_paragraph(doc, "{{판단흐름}}")
    add_paragraph(doc, "Outcome: {{판정결과}}", bold=True)
    add_h1(doc, "5. Conclusion of MDCG 2020-3 Assessment")
    add_paragraph(doc, "{{판정요약}}")
    add_h1(doc, "6. Technical Documentation Updates (MDR Annex II / III)")
    add_bullet(doc, "Annex II – Section 1 (Device Description and Specification): Updated where applicable.")
    add_bullet(doc, "Annex II – BOM and component list: Updated.")
    add_bullet(doc, "Annex II – Labelling and IFU: Updated where the changed item is referenced.")
    add_h1(doc, "7. References")
    add_bullet(doc, "Regulation (EU) 2017/745 (MDR)")
    add_bullet(doc, "MDCG 2020-3 Rev.1")
    add_bullet(doc, "Notified Body Certificate No. {{EU_인증서번호}}")
    return doc


def _kr():
    doc = init_document()
    add_paragraph(doc, "국내 의료기기 변경사항 검토서 (경미한 변경사항 판단)", bold=True)
    add_spacer(doc)
    add_paragraph(doc, "1) 대상 제품", bold=True)
    add_paragraph(doc, "- {{국내허가번호}} ({{모델명}})")
    add_paragraph(doc, "- 검토일: {{평가일}}")
    add_spacer(doc)
    add_paragraph(doc, "2) 변경 사항", bold=True)
    add_paragraph(doc, "(1) 변경 전 항목", bold=True)
    add_paragraph(doc, "{{구성품명}}", bold=True)
    add_info_table(doc, [("변경 전", "{{변경전}}")])
    add_spacer(doc)
    add_paragraph(doc, "(2) 변경 후 항목", bold=True)
    add_paragraph(doc, "{{구성품명}}", bold=True)
    add_info_table(doc, [("변경 후", "{{변경후}}")])
    add_spacer(doc)
    add_paragraph(doc, "{{변경상세}}")
    add_spacer(doc)
    add_paragraph(doc, "(3) 변경 사유", bold=True)
    add_paragraph(doc, "{{변경사유}}")
    add_page_break(doc)
    add_paragraph(doc, "{{가이던스명}}에 따른 변경대상 판단 결과", bold=True)
    add_spacer(doc)
    add_paragraph(doc, "- 판단 흐름", bold=True)
    add_paragraph(doc, "{{판단흐름}}")
    add_spacer(doc)
    add_paragraph(doc, "결론: {{판정결과}}")
    add_spacer(doc)
    add_paragraph(doc, "{{판정요약}}", bold=True)
    add_spacer(doc)
    add_paragraph(doc, "서명", bold=True)
    add_info_table(doc, [("작성자", "{{작성자}}"), ("검토자", "{{검토자}}"), ("승인자", "{{승인자}}")])
    return doc


_BUILDERS = {"FDA": _fda, "HC": _hc, "EU": _eu, "KR": _kr}


def default_template_bytes(country: str) -> bytes:
    buf = BytesIO()
    _BUILDERS[country]().save(buf)
    return buf.getvalue()
